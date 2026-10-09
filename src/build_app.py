#!/usr/bin/env python3
"""Construye la app Excel "Gaudí · Visualizador de reformas".

Genera en dist/:
  Gaudi_Reformas.xlsm        App completa (macros): generación automática con IA en Excel de escritorio.
  Gaudi_Reformas_Movil.xlsx  Sin macros, para móvil/tablet: modo asistido (prompt + Gemini/ChatGPT).

Uso:  python3 src/build_app.py
"""

from __future__ import annotations

import io
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from pathlib import Path

import xlsxwriter

import listas as L
from vbaproject import Module, build_vba_project

SRC = Path(__file__).resolve().parent
ROOT = SRC.parent
DIST = ROOT / "dist"
VBA_DIR = SRC / "vba"

XLSM = DIST / "Gaudi_Reformas.xlsm"
XLSX_MOVIL = DIST / "Gaudi_Reformas_Movil.xlsx"

# Paleta inspirada en los trencadís de Gaudí
TEAL = "#1D4E5A"
TERRACOTA = "#C8553D"
MOSTAZA = "#E9A23B"
ARENA = "#F7F1E6"
CREMA = "#FFF6D8"
BORDE_CREMA = "#D9B26F"
TINTA = "#2E2E2E"
GRIS = "#7A7A7A"
NEUTRO = "#6B5B4B"
BLANCO = "#FFFFFF"

FUENTE = "Arial"

HOJAS = [  # (nombre visible, codeName VBA)
    ("INICIO", "shInicio"),
    ("PROYECTO", "shProyecto"),
    ("RESULTADO", "shResultado"),
    ("HISTORIAL", "shHistorial"),
    ("AJUSTES", "shAjustes"),
    ("LISTAS", "shListas"),
    ("CALC", "shCalc"),
]

PROTECCION = {
    "objects": True,  # permite insertar/mover imágenes (fotos en el móvil)
    "select_locked_cells": True,
    "select_unlocked_cells": True,
}

EJEMPLO = {  # proyecto de ejemplo con valores realistas
    "ProyNombre": "Ejemplo · Salón calle Mallorca",
    "ProyEstancia": "Salón-comedor",
    "ProyTipo": "Reforma integral",
    "ProyEstilo": "Nórdico",
    "ProyIntensidad": "Media",
    "RefSuelo": "Parquet de roble natural",
    "RefParedes": "Pintura lisa",
    "RefColor": "Blanco roto",
    "RefTecho": "Liso blanco",
    "RefCarpinteria": "Puertas lacadas en blanco",
    "RefCocina": "Mantener / no aplica",
    "RefEncimera": "Mantener / no aplica",
    "RefBano": "Mantener / no aplica",
    "RefIluminacion": "Lámparas colgantes de diseño",
    "RefMobiliario": "Amueblar según el estilo",
    "RefNotas": "Sofá en tonos claros, alfombra de yute y plantas naturales",
    "FmtProporcion": "Igual que la foto original",
    "FmtCalidad": "Alta (2K)",
    "FmtEstilo": "Foto realista",
    "FmtLuz": "Luz natural de día",
    "EstConservar": "",
}


# =============================================================================
# VBA: convierte el código a ASCII puro (los acentos de los textos pasan a ChrW)
# =============================================================================

def vba_ascii(code: str, nombre: str) -> str:
    salida = []
    for num, linea in enumerate(code.splitlines(), 1):
        res: list[str] = []
        i, n = 0, len(linea)
        en_cadena = False
        partes: list[str] = []
        tramo: list[str] = []
        tiene_unicode = False
        while i < n:
            ch = linea[i]
            if not en_cadena:
                if ch == '"':
                    en_cadena, partes, tramo = True, [], []
                elif ch == "'":
                    comentario = unicodedata.normalize("NFKD", linea[i:]).encode("ascii", "ignore").decode()
                    res.append(comentario)
                    break
                else:
                    if ord(ch) > 127:
                        raise ValueError(f"{nombre}:{num}: carácter no ASCII fuera de una cadena")
                    res.append(ch)
                i += 1
                continue
            if ch == '"':
                if i + 1 < n and linea[i + 1] == '"':
                    tramo.append('""')
                    i += 2
                    continue
                if tramo or not partes:
                    partes.append('"' + "".join(tramo) + '"')
                res.append(partes[0] if len(partes) == 1 else "(" + " & ".join(partes) + ")")
                en_cadena = False
                i += 1
                continue
            if ord(ch) > 127:
                if ord(ch) > 0xFFFF:
                    raise ValueError(f"{nombre}:{num}: carácter fuera de BMP")
                tiene_unicode = True
                if tramo:
                    partes.append('"' + "".join(tramo) + '"')
                    tramo = []
                partes.append(f"ChrW({ord(ch)})")
            else:
                tramo.append(ch)
            i += 1
        if en_cadena:
            raise ValueError(f"{nombre}:{num}: cadena sin cerrar")
        nueva = "".join(res)
        if tiene_unicode and re.search(r"\b(Const|Optional)\b", linea):
            raise ValueError(f"{nombre}:{num}: acentos en Const/Optional no permitidos")
        if len(nueva) > 1000:
            raise ValueError(f"{nombre}:{num}: línea demasiado larga para VBA")
        salida.append(nueva)
    return "\n".join(salida) + "\n"


def proyecto_vba() -> bytes:
    modulos = [Module("ThisWorkbook", "workbook", "")]
    modulos += [Module(code, "sheet", "") for _, code in HOJAS]
    for nombre in ("modGaudi", "modIA", "modUtil"):
        codigo = (VBA_DIR / f"{nombre}.bas").read_text(encoding="utf-8")
        modulos.append(Module(nombre, "module", vba_ascii(codigo, nombre)))
    return build_vba_project(modulos)


# =============================================================================
# Libro
# =============================================================================

class Formatos:
    def __init__(self, wb: xlsxwriter.Workbook):
        self.wb = wb
        self._cache: dict[tuple, object] = {}

    def __call__(self, **props):
        clave = tuple(sorted(props.items()))
        if clave not in self._cache:
            base = {"font_name": FUENTE, "font_size": 11, "font_color": TINTA, "valign": "vcenter"}
            base.update(props)
            self._cache[clave] = self.wb.add_format(base)
        return self._cache[clave]


class App:
    def __init__(self, ruta: Path, macros: bool, cache: dict[str, object]):
        self.macros = macros
        self.cache = cache
        self.wb = xlsxwriter.Workbook(str(ruta), {
            "default_format_properties": {"font_name": FUENTE, "font_size": 11},
        })
        self.f = Formatos(self.wb)
        self.ws = {}
        for nombre, code in HOJAS:
            ws = self.wb.add_worksheet(nombre)
            if macros:
                ws.set_vba_name(code)
            self.ws[nombre] = ws
        if macros:
            self.wb.set_vba_name("ThisWorkbook")
        self.nombres: dict[str, str] = {}
        self.botones = 0

    # ---------------------------------------------------------------- utilidades
    def nombre(self, nombre: str, hoja: str, celda: str):
        col, fila = re.match(r"([A-Z]+)(\d+)", celda).groups()
        ref = f"={hoja}!${col}${fila}"
        if ":" in celda:
            a, b = celda.split(":")
            ca, fa = re.match(r"([A-Z]+)(\d+)", a).groups()
            cb, fb = re.match(r"([A-Z]+)(\d+)", b).groups()
            ref = f"={hoja}!${ca}${fa}:${cb}${fb}"
        self.wb.define_name(nombre, ref)
        self.nombres[nombre] = f"{hoja}!{celda.split(':')[0]}"

    def formula(self, hoja: str, celda: str, formula: str, fmt):
        valor = self.cache.get(f"{hoja}!{celda}", "")
        self.ws[hoja].write_formula(celda, formula, fmt, valor)

    def barra_titulo(self, ws, rango: str, texto: str, alto=40, tam=20):
        ws.merge_range(rango, texto, self.f(bold=True, font_size=tam, font_color=BLANCO, bg_color=TEAL,
                                            indent=1))
        fila = int(re.search(r"(\d+)", rango).group(1)) - 1
        ws.set_row(fila, alto)

    def seccion(self, ws, rango: str, texto: str, color=TEAL):
        ws.merge_range(rango, texto, self.f(bold=True, font_size=11, font_color=BLANCO, bg_color=color,
                                            indent=1))
        fila = int(re.search(r"(\d+)", rango).group(1)) - 1
        ws.set_row(fila, 24)

    def boton(self, hoja: str, celda: str, texto: str, macro: str, color=TEAL, ancho=270, alto=40,
              dx=0, dy=4, texto_oscuro=False):
        if not self.macros:
            return
        self.botones += 1
        self.ws[hoja].insert_textbox(celda, texto, {
            "width": ancho, "height": alto, "x_offset": dx, "y_offset": dy,
            "font": {"name": FUENTE, "bold": True, "size": 10, "color": TINTA if texto_oscuro else BLANCO},
            "align": {"vertical": "middle", "horizontal": "center", "text": "center"},
            "fill": {"color": color},
            "line": {"none": True},
            "description": f"btn:{macro}",
            "object_position": 3,
        })

    def imprimir(self, ws, area: str, horizontal=False, paginas_alto=1):
        ws.set_paper(9)  # A4
        if horizontal:
            ws.set_landscape()
        ws.set_margins(left=0.4, right=0.4, top=0.5, bottom=0.5)
        ws.center_horizontally()
        ws.print_area(area)
        ws.fit_to_pages(1, paginas_alto)

    def desplegable(self, hoja: str, celda: str, lista: str, titulo: str, mensaje: str, estricto=True):
        opciones = {
            "validate": "list", "source": f"=L_{lista}",
            "input_title": titulo[:32], "input_message": mensaje[:255],
        }
        if estricto:
            opciones.update({"error_title": "Opción no válida",
                             "error_message": "Elige una opción de la lista desplegable."})
        else:
            opciones.update({"error_type": "information", "error_title": "Valor personalizado",
                             "error_message": "Valor fuera de la lista. Se usará tal cual."})
        self.ws[hoja].data_validation(celda, opciones)

    # ---------------------------------------------------------------- hojas
    def listas(self):
        ws = self.ws["LISTAS"]
        cab = self.f(bold=True, bg_color=TEAL, font_color=BLANCO)
        txt = self.f(font_size=9)
        col = 0
        for nombre, (filas, cabeceras) in L.TODAS.items():
            ancho = len(cabeceras)
            for k, c in enumerate(cabeceras):
                ws.write(0, col + k, c, cab)
            for r, fila in enumerate(filas, start=1):
                for k, v in enumerate(fila):
                    ws.write_string(r, col + k, v, txt)
            ultima = len(filas) + 1
            letra = xlsxwriter.utility.xl_col_to_name(col)
            letra_fin = xlsxwriter.utility.xl_col_to_name(col + ancho - 1)
            self.wb.define_name(f"L_{nombre}", f"=LISTAS!${letra}$2:${letra}${ultima}")
            self.wb.define_name(f"T_{nombre}", f"=LISTAS!${letra}$2:${letra_fin}${ultima}")
            ws.set_column(col, col, 30)
            if ancho > 1:
                ws.set_column(col + 1, col + ancho - 1, 40)
            col += ancho + 1
        ws.protect("", PROTECCION)
        ws.hide()

    def calc(self):
        """Hoja oculta que traduce cada elección a un fragmento en inglés y monta el prompt."""
        ws = self.ws["CALC"]
        ws.set_column("A:A", 22)
        ws.set_column("B:B", 100)
        ws.set_column("C:C", 50)
        cab = self.f(bold=True, bg_color=TEAL, font_color=BLANCO)
        ws.write_row("A1", ["Clave", "Valor (fórmula)", "Descripción"], cab)
        txt = self.f(font_size=9, text_wrap=True, valign="top")

        filas = [
            ("EN_Estancia", 'IFERROR(VLOOKUP(ProyEstancia,T_Estancia,2,FALSE),"room")', "Estancia en inglés"),
            ("EN_Tipo", 'IFERROR(VLOOKUP(ProyTipo,T_Tipo,2,FALSE),"renovation")', "Tipo de reforma"),
            ("EN_Estilo", 'IFERROR(VLOOKUP(ProyEstilo,T_Estilo,2,FALSE),"modern")', "Estilo decorativo"),
            ("EN_Intensidad", 'IFERROR(VLOOKUP(ProyIntensidad,T_Intensidad,2,FALSE),"")', "Intensidad"),
            ("EN_Suelo", 'IFERROR(VLOOKUP(RefSuelo,T_Suelo,2,FALSE),"keep the existing floor")', "Suelo"),
            ("EN_Paredes", 'IFERROR(VLOOKUP(RefParedes,T_Paredes,2,FALSE),"keep the existing wall finish")',
             "Acabado de paredes"),
            ("EN_Color", 'IFERROR(VLOOKUP(RefColor,T_Color,2,FALSE),"keep the existing colors")', "Color"),
            ("EN_Techo", 'IFERROR(VLOOKUP(RefTecho,T_Techo,2,FALSE),"keep the existing ceiling")', "Techo"),
            ("EN_Carpinteria", 'IFERROR(VLOOKUP(RefCarpinteria,T_Carpinteria,2,FALSE),'
                               '"keep the existing doors and windows")', "Puertas y ventanas"),
            ("EN_Cocina", 'IFERROR(VLOOKUP(RefCocina,T_Cocina,2,FALSE),"-")', "Cocina ('-' = no se menciona)"),
            ("EN_Encimera", 'IFERROR(VLOOKUP(RefEncimera,T_Encimera,2,FALSE),"-")', "Encimera"),
            ("EN_Bano", 'IFERROR(VLOOKUP(RefBano,T_Bano,2,FALSE),"-")', "Baño"),
            ("EN_Iluminacion", 'IFERROR(VLOOKUP(RefIluminacion,T_Iluminacion,2,FALSE),'
                               '"keep the existing light fittings")', "Lámparas"),
            ("EN_Mobiliario", 'IFERROR(VLOOKUP(RefMobiliario,T_Mobiliario,2,FALSE),"keep the existing furniture")',
             "Mobiliario"),
            ("EN_FmtEstilo", 'IFERROR(VLOOKUP(FmtEstilo,T_EstiloImagen,2,FALSE),"photorealistic interior photograph")',
             "Estilo de imagen"),
            ("EN_FmtLuz", 'IFERROR(VLOOKUP(FmtLuz,T_Luz,2,FALSE),"keep the original lighting")', "Luz"),
            ("CodProporcion", 'IFERROR(VLOOKUP(FmtProporcion,T_Proporcion,2,FALSE),"auto")',
             "Proporción para la API (auto = la de la foto)"),
            ("EN_Proporcion", 'IFERROR(VLOOKUP(FmtProporcion,T_Proporcion,3,FALSE),"")', "Texto de proporción"),
            ("CodResolucion", 'IFERROR(VLOOKUP(FmtCalidad,T_Calidad,2,FALSE),"1K")', "imageSize de Gemini"),
            ("CodCalidadOpenAI", 'IFERROR(VLOOKUP(FmtCalidad,T_Calidad,3,FALSE),"medium")', "quality de OpenAI"),
        ]
        celdas: dict[str, str] = {}
        r = 2
        for clave, formula, desc in filas:
            celdas[clave] = f"CALC!$B${r}"
            ws.write(f"A{r}", clave, txt)
            self.formula("CALC", f"B{r}", "=" + formula, txt)
            ws.write(f"C{r}", desc, txt)
            if clave.startswith("Cod"):
                self.nombre(clave, "CALC", f"B{r}")
            r += 1

        c = celdas
        estructura = (
            "STRICT RULE - KEEP THE ORIGINAL ARCHITECTURE: this is the same real space after the renovation. "
            "Keep exactly the same camera position, viewpoint, height, angle, lens, perspective, framing and "
            "vanishing lines as in the original photo. Do not move, add, remove, resize or reshape any wall, "
            "partition, window, door, opening, column, beam, arch, staircase, radiator or built-in element. "
            "Keep the room dimensions, ceiling height, layout and proportions identical, so the new image "
            "overlays perfectly on the original one. Only surfaces, finishes, colors, fixtures, light fittings "
            "and furniture may change, as specified below."
        )
        prompt_filas = [
            ("P_Intro",
             f'"Interior renovation visualization. Edit the attached photo (space: "&{c["EN_Estancia"]}'
             f'&") to show how it will look after this renovation: "&{c["EN_Tipo"]}&"."',
             "Introducción"),
            ("P_Estructura",
             f'"{estructura}"&IF(LEN(TRIM(EstConservar))>0," Also keep exactly as they are: "'
             f'&TRIM(EstConservar)&".","")',
             "Bloqueo de estructura (siempre presente)"),
            ("P_Estilo",
             f'"Interior design style: "&{c["EN_Estilo"]}&". "&{c["EN_Intensidad"]}',
             "Estilo e intensidad"),
            ("P_Cambios",
             f'"Changes: floor: "&{c["EN_Suelo"]}&". Wall finish: "&{c["EN_Paredes"]}'
             f'&". Wall color: "&{c["EN_Color"]}&". Ceiling: "&{c["EN_Techo"]}'
             f'&". Doors and windows: "&{c["EN_Carpinteria"]}&"."'
             f'&IF({c["EN_Cocina"]}="-",""," Kitchen cabinets: "&{c["EN_Cocina"]}&".")'
             f'&IF({c["EN_Encimera"]}="-",""," Countertop: "&{c["EN_Encimera"]}&".")'
             f'&IF({c["EN_Bano"]}="-",""," Bathroom: "&{c["EN_Bano"]}&".")'
             f'&" Light fittings: "&{c["EN_Iluminacion"]}&". Furniture and decor: "&{c["EN_Mobiliario"]}&"."',
             "Cambios por elemento"),
            ("P_Notas",
             'IF(LEN(TRIM(RefNotas))>0,"Additional client requests (written in Spanish; follow them as long as '
             'they do not change the architecture): "&TRIM(RefNotas)&".","")',
             "Notas del usuario"),
            ("P_Salida",
             f'"Output: "&{c["EN_FmtEstilo"]}&". Lighting: "&{c["EN_FmtLuz"]}&". "&{c["EN_Proporcion"]}'
             f'&" Realistic materials and textures, correct real-world scale, coherent shadows and reflections. '
             f'No text, labels, watermarks or people."',
             "Formato de salida"),
        ]
        r += 1
        for clave, formula, desc in prompt_filas:
            celdas[clave] = f"CALC!$B${r}"
            ws.write(f"A{r}", clave, txt)
            self.formula("CALC", f"B{r}", "=" + formula, txt)
            ws.write(f"C{r}", desc, txt)
            r += 1
        self.calc_celdas = celdas
        ws.protect("", PROTECCION)
        ws.hide()

    def inicio(self):
        ws = self.ws["INICIO"]
        f = self.f
        ws.hide_gridlines(2)
        ws.set_column("A:A", 1.5)
        ws.set_column("B:B", 5)
        ws.set_column("C:C", 62)
        ws.set_column("D:D", 1.5)
        ws.set_default_row(20)
        ws.set_tab_color(TEAL)
        texto = f(text_wrap=True, valign="top")
        num = f(bold=True, font_color=BLANCO, bg_color=TERRACOTA, align="center", valign="top")
        letra = f(bold=True, font_color=BLANCO, bg_color=TEAL, align="center", valign="top")

        self.barra_titulo(ws, "B1:C1", "GAUDÍ", alto=46, tam=26)
        ws.merge_range("B2:C2", "Visualizador de reformas con IA",
                       f(bold=True, font_size=13, font_color=BLANCO, bg_color=TERRACOTA, indent=1))
        ws.set_row(1, 26)
        ws.set_row(2, 8)
        ws.merge_range("B4:C4",
                       "Sube la foto de un piso o de una estancia, define el formato y las reformas, y obtén una "
                       "imagen de cómo quedará. La estructura original (paredes, ventanas, puertas, techos y "
                       "perspectiva) se mantiene siempre.", f(text_wrap=True, valign="top", font_size=12))
        ws.set_row(3, 66)

        self.seccion(ws, "B6:C6", "CÓMO SE USA")
        pasos = [
            "Ve a PROYECTO y carga la foto original (en el móvil la adjuntarás directamente en Gemini o ChatGPT).",
            "Elige estancia, estilo, reformas y formato en los desplegables de color crema.",
            "Pulsa GENERAR IMAGEN: verás el ANTES y el DESPUÉS en la hoja RESULTADO.",
        ]
        for k, p in enumerate(pasos):
            ws.write(6 + k, 1, k + 1, num)
            ws.write(6 + k, 2, p, texto)
            ws.set_row(6 + k, 34)

        self.seccion(ws, "B11:C11", "DOS FORMAS DE GENERAR LA IMAGEN")
        if self.macros:
            modos = [
                ("A", "Automática (Excel de escritorio con macros): la app envía la foto a Google Gemini u OpenAI "
                      "con tu clave API y coloca el resultado en RESULTADO. Windows: completa. Mac: experimental."),
                ("B", "Asistida (móvil, tablet o sin clave): copia el PROMPT de la sección 7 de PROYECTO, abre "
                      "Gemini o ChatGPT, adjunta la foto y pega el texto. Funciona con una cuenta gratuita."),
            ]
        else:
            modos = [
                ("B", "Asistida (esta versión para móvil y tablet): copia el PROMPT de la sección 7 de PROYECTO, "
                      "abre Gemini o ChatGPT, adjunta la foto y pega el texto. Funciona con una cuenta gratuita."),
                ("A", "Automática: usa el archivo Gaudi_Reformas.xlsm en Excel de escritorio (lleva macros y "
                      "conecta con la IA usando tu clave API)."),
            ]
        for k, (l, t) in enumerate(modos):
            ws.write(11 + k, 1, l, letra)
            ws.write(11 + k, 2, t, texto)
            ws.set_row(11 + k, 50)

        self.seccion(ws, "B15:C15", "PRIMERA VEZ EN EL ORDENADOR" if self.macros else "EN EL MÓVIL")
        if self.macros:
            primera = [
                "Si Excel avisa de que las macros están bloqueadas: cierra Excel, clic derecho en el archivo "
                "› Propiedades › marca «Desbloquear» › Aceptar. Ábrelo de nuevo y pulsa «Habilitar contenido».",
                "En AJUSTES pulsa GUARDAR CLAVE GEMINI (la clave se crea gratis en aistudio.google.com/apikey).",
                "Pulsa PROBAR CONEXIÓN y ya puedes generar desde PROYECTO.",
            ]
        else:
            primera = [
                "Abre PROYECTO y elige las opciones en los desplegables (toca la celda y luego la flecha).",
                "Baja a la sección 7, mantén pulsada la celda del PROMPT y elige Copiar.",
                "Toca «Abrir Gemini» o «Abrir ChatGPT», adjunta tu foto, pega el texto y envía.",
            ]
        for k, p in enumerate(primera):
            ws.write(15 + k, 1, k + 1, num)
            ws.write(15 + k, 2, p, texto)
            ws.set_row(15 + k, 50 if k == 0 else 34)

        self.seccion(ws, "B20:C20", "IR A…")
        enlace = f(font_color=TEAL, bold=True, underline=1, font_size=12)
        destinos = [
            ("internal:'PROYECTO'!C10", "→  PROYECTO · configurar la reforma"),
            ("internal:'RESULTADO'!A1", "→  RESULTADO · antes y después"),
            ("internal:'HISTORIAL'!A1", "→  HISTORIAL · imágenes generadas"),
            ("internal:'AJUSTES'!C6", "→  AJUSTES · motor de IA y clave API"),
        ]
        for k, (url, t) in enumerate(destinos):
            ws.merge_range(20 + k, 1, 20 + k, 2, "", enlace)
            ws.write_url(20 + k, 1, url, enlace, t)
            ws.set_row(20 + k, 26)

        self.seccion(ws, "B26:C26", "LEYENDA")
        ws.write(26, 1, "", f(bg_color=CREMA, border=1, border_color=BORDE_CREMA))
        ws.write(26, 2, "Celda editable (desplegable o texto): solo hay que tocar estas.", texto)
        ws.write(27, 1, "", f(bg_color=BLANCO, border=1, border_color="#BFBFBF"))
        ws.write(27, 2, "Celda calculada o rellenada por la app: no se edita.", texto)
        ws.write(28, 1, "", f(bg_color=TERRACOTA))
        ws.write(28, 2, "Botón de acción (solo en Excel de escritorio con macros).", texto)
        for k in range(26, 29):
            ws.set_row(k, 22)

        self.seccion(ws, "B31:C31", "PRIVACIDAD Y AVISO")
        ws.merge_range("B32:C32",
                       "La foto y las especificaciones solo se envían al proveedor de IA elegido (Google u OpenAI) "
                       "cuando generas una imagen. La clave API se guarda en tu equipo, no dentro del archivo. Las "
                       "imágenes son orientativas: no sustituyen a un proyecto técnico ni a un presupuesto.",
                       f(text_wrap=True, valign="top", font_size=10, font_color=GRIS))
        ws.set_row(31, 64)
        self.imprimir(ws, "B1:C32")
        ws.protect("", PROTECCION)
        ws.activate()

    def proyecto(self):
        ws = self.ws["PROYECTO"]
        f = self.f
        ws.hide_gridlines(2)
        ws.set_tab_color(TERRACOTA)
        ws.set_column("A:A", 1.5)
        ws.set_column("B:B", 24)
        ws.set_column("C:C", 38)
        ws.set_column("D:D", 2.5)
        ws.set_column("E:J", 12.5)
        ws.set_column("K:K", 2)
        ws.set_default_row(22)

        etiqueta = f(bold=True, font_size=10, bg_color=ARENA, border=1, border_color="#E6DCCB", indent=1,
                     text_wrap=True)
        entrada = f(bg_color=CREMA, border=1, border_color=BORDE_CREMA, locked=False, text_wrap=True, indent=1)
        calculada = f(border=1, border_color="#D9D9D9", text_wrap=True, font_size=10)
        pequena = f(font_size=8, font_color=GRIS, border=1, border_color="#D9D9D9", text_wrap=True)
        enlace_nav = f(font_color=TEAL, bold=True, underline=1, font_size=10)

        self.barra_titulo(ws, "B1:C1", "GAUDÍ · Proyecto de reforma")
        if self.macros:
            ws.merge_range("E1:J1", "", f(bg_color=TEAL))
        ws.write_url("B2", "internal:'INICIO'!A1", enlace_nav, "← Inicio")
        ws.write_url("C2", "internal:'RESULTADO'!A1", f(font_color=TEAL, bold=True, underline=1, font_size=10,
                                                      align="right"), "Ver RESULTADO →")
        ws.merge_range("B3:C3", "Rellena de arriba abajo las celdas color crema (toca la celda y elige en la lista).",
                       f(font_size=9, font_color=GRIS, italic=True, text_wrap=True))

        fila = 4  # fila Excel (1-based) actual

        def sec(texto):
            nonlocal fila
            self.seccion(ws, f"B{fila}:C{fila}", texto)
            fila += 1

        def campo(etq, nombre, lista=None, msg="", alto=None, estricto=True, fmt=None):
            nonlocal fila
            ws.write(f"B{fila}", etq, etiqueta)
            valor = EJEMPLO.get(nombre, "")
            ws.write_string(f"C{fila}", valor, fmt or entrada)
            self.nombre(nombre, "PROYECTO", f"C{fila}")
            if lista:
                self.desplegable("PROYECTO", f"C{fila}", lista, etq, msg, estricto)
            elif msg:
                ws.data_validation(f"C{fila}", {"validate": "any", "input_title": etq[:32],
                                                "input_message": msg[:255]})
            if alto:
                ws.set_row(fila - 1, alto)
            fila += 1

        def hueco(alto=8):
            nonlocal fila
            ws.set_row(fila - 1, alto)
            fila += 1

        # 1 · Foto
        fila_foto = fila
        sec("1 · FOTO ORIGINAL")
        ws.write(f"B{fila}", "Foto", etiqueta)
        ws.write(f"C{fila}", "(ninguna: pulsa CARGAR FOTO)" if self.macros else "(en el móvil: adjúntala en Gemini)",
                 calculada)
        self.nombre("FotoNombre", "PROYECTO", f"C{fila}")
        fila += 1
        ws.write(f"B{fila}", "Ruta del archivo", etiqueta)
        ws.write(f"C{fila}", "", pequena)
        self.nombre("FotoRuta", "PROYECTO", f"C{fila}")
        fila += 1
        hueco()

        # 2 · Proyecto
        sec("2 · PROYECTO Y ESTANCIA")
        campo("Nombre del proyecto", "ProyNombre", msg="Texto libre. Se usa para nombrar las imágenes.")
        campo("Estancia", "ProyEstancia", "Estancia", "¿Qué se ve en la foto?")
        campo("Tipo de reforma", "ProyTipo", "Tipo", "Alcance de la reforma.")
        campo("Estilo decorativo", "ProyEstilo", "Estilo", "Estilo de interiorismo del resultado.")
        campo("Intensidad del cambio", "ProyIntensidad", "Intensidad",
              "Cuánto cambian acabados y muebles. La arquitectura nunca cambia.")
        hueco()

        # 3 · Reformas
        sec("3 · REFORMAS POR ELEMENTO")
        campo("Suelo", "RefSuelo", "Suelo", "Nuevo suelo o «Mantener actual».")
        campo("Paredes · acabado", "RefParedes", "Paredes", "Revestimiento de las paredes.")
        campo("Paredes · color", "RefColor", "Color", "Color principal de paredes.")
        campo("Techo", "RefTecho", "Techo", "Acabado del techo (la altura no cambia).")
        campo("Puertas y ventanas", "RefCarpinteria", "Carpinteria",
              "Acabado de carpinterías. Huecos y tamaños se mantienen.")
        campo("Cocina · muebles", "RefCocina", "Cocina", "Solo si en la foto hay cocina.")
        campo("Cocina · encimera", "RefEncimera", "Encimera", "Solo si en la foto hay cocina.")
        campo("Baño", "RefBano", "Bano", "Solo si en la foto hay baño. Los sanitarios quedan en su sitio.")
        campo("Iluminación", "RefIluminacion", "Iluminacion", "Lámparas y puntos de luz.")
        campo("Mobiliario y decoración", "RefMobiliario", "Mobiliario", "Qué hacer con los muebles.")
        campo("Notas adicionales", "RefNotas", msg="Texto libre (opcional). Ej.: sofá verde, estantería a medida.",
              alto=44)
        hueco()

        # 4 · Formato
        sec("4 · FORMATO DE LA IMAGEN")
        campo("Proporción", "FmtProporcion", "Proporcion",
              "Recomendado: igual que la foto, para no recortar ni inventar partes de la estancia.")
        campo("Calidad / resolución", "FmtCalidad", "Calidad", "Más calidad = más lenta (y más cara con API).")
        campo("Estilo de imagen", "FmtEstilo", "EstiloImagen", "Foto realista, render, boceto…")
        campo("Luz de la escena", "FmtLuz", "Luz", "Ambiente de luz del resultado.")
        hueco()

        # 5 · Estructura
        sec("5 · ESTRUCTURA ORIGINAL (SIEMPRE SE CONSERVA)")
        ws.merge_range(f"B{fila}:C{fila}",
                       "🔒 Se mantienen intactos: paredes y tabiques, ventanas, puertas y huecos, techo y altura, "
                       "vigas, pilares, escaleras, dimensiones y la perspectiva exacta de la foto.",
                       f(text_wrap=True, font_size=10, bg_color="#E8F0F1", border=1, border_color="#C5D6D9",
                         valign="top"))
        ws.set_row(fila - 1, 46)
        fila += 1
        campo("Conservar también", "EstConservar",
              msg="Elementos concretos que no deben cambiar. Ej.: la chimenea, el radiador de hierro.")
        hueco()

        # 6 · Resumen
        fila_generar = fila
        sec("6 · RESUMEN Y ESTADO")
        ws.write(f"B{fila}", "Estado", etiqueta)
        estado_ini = ("Listo. Carga una foto para empezar." if self.macros
                      else "Versión móvil: usa la sección 7 para generar con Gemini o ChatGPT.")
        ws.write(f"C{fila}", estado_ini, f(font_size=10, border=1, border_color="#D9D9D9", text_wrap=True,
                                          font_color=TEAL, bold=True))
        self.nombre("EstadoApp", "PROYECTO", f"C{fila}")
        ws.set_row(fila - 1, 34)
        fila += 1
        c = self.calc_celdas
        resumen = (
            '="Estancia: "&ProyEstancia&"  ·  "&ProyTipo&CHAR(10)'
            '&"Estilo: "&ProyEstilo&" (intensidad "&LOWER(ProyIntensidad)&")"&CHAR(10)'
            '&"Suelo: "&RefSuelo&"  ·  Paredes: "&RefParedes&" / "&RefColor&CHAR(10)'
            '&"Techo: "&RefTecho&"  ·  Carpintería: "&RefCarpinteria&CHAR(10)'
            f'&IF(AND({c["EN_Cocina"]}="-",{c["EN_Encimera"]}="-"),"",'
            '"Cocina: "&RefCocina&"  ·  Encimera: "&RefEncimera&CHAR(10))'
            f'&IF({c["EN_Bano"]}="-","","Baño: "&RefBano&CHAR(10))'
            '&"Iluminación: "&RefIluminacion&"  ·  Mobiliario: "&RefMobiliario&CHAR(10)'
            '&"Imagen: "&FmtProporcion&"  ·  "&FmtCalidad&"  ·  "&FmtEstilo&"  ·  "&FmtLuz'
            '&IF(LEN(TRIM(RefNotas))>0,CHAR(10)&"Notas: "&TRIM(RefNotas),"")'
            '&IF(LEN(TRIM(EstConservar))>0,CHAR(10)&"Conservar también: "&TRIM(EstConservar),"")'
        )
        fmt_res = f(text_wrap=True, font_size=10, valign="top", border=1, border_color="#D9D9D9", bg_color=ARENA)
        ws.merge_range(f"B{fila}:C{fila}", "", fmt_res)
        self.formula("PROYECTO", f"B{fila}", resumen, fmt_res)
        self.nombre("ResumenES", "PROYECTO", f"B{fila}")
        ws.set_row(fila - 1, 128)
        fila += 1
        hueco()

        # 7 · Modo móvil
        sec("7 · MODO MÓVIL / SIN CLAVE API")
        ws.merge_range(f"B{fila}:C{fila}",
                       "1) Mantén pulsada la celda del PROMPT (abajo) y elige Copiar.  2) Abre Gemini o ChatGPT.  "
                       "3) Adjunta la foto original, pega el texto y envía.  4) Guarda la imagen que te devuelva"
                       + (" (en el ordenador puedes traerla con IMPORTAR RESULTADO)." if self.macros else "."),
                       f(text_wrap=True, font_size=10, valign="top"))
        ws.set_row(fila - 1, 58)
        fila += 1
        prompt = "=" + "&\" \"&".join([
            c["P_Intro"], c["P_Estructura"], c["P_Estilo"], c["P_Cambios"]]) + \
            f'&IF({c["P_Notas"]}="",""," "&{c["P_Notas"]})&" "&{c["P_Salida"]}'
        fmt_prompt = f(text_wrap=True, font_size=9, valign="top", border=2, border_color=TERRACOTA, bg_color=BLANCO)
        ws.merge_range(f"B{fila}:C{fila}", "", fmt_prompt)
        self.formula("PROYECTO", f"B{fila}", prompt, fmt_prompt)
        self.nombre("PromptIA", "PROYECTO", f"B{fila}")
        largo = len(str(self.cache.get(f"PROYECTO!B{fila}", ""))) or 2200
        lineas = largo / 92 + 3
        ws.set_row(fila - 1, min(409, max(150, lineas * 11)))
        self.fila_prompt = fila
        fila += 1
        enlace = f(font_color=BLANCO, bold=True, bg_color=TEAL, align="center", font_size=11)
        ws.write_url(f"B{fila}", "https://gemini.google.com/app", enlace, "Abrir Gemini →")
        ws.write_url(f"C{fila}", "https://chatgpt.com/", f(font_color=BLANCO, bold=True, bg_color="#10A37F",
                                                            align="center", font_size=11), "Abrir ChatGPT →")
        ws.set_row(fila - 1, 30)
        fila += 1
        ws.merge_range(f"B{fila}:C{fila}",
                       "El prompt se escribe en inglés porque los modelos de imagen lo siguen con más precisión. "
                       "Se actualiza solo al cambiar cualquier opción.",
                       f(font_size=8, italic=True, font_color=GRIS, text_wrap=True))
        ws.set_row(fila - 1, 26)

        # Panel de escritorio (columna derecha): vista previa y botones con macros
        caja_ini, caja_fin = fila_foto + 1, fila_foto + 16
        self.nombre("CajaOriginal", "PROYECTO", f"E{caja_ini}:J{caja_fin}")
        if self.macros:
            self.seccion(ws, f"E{fila_foto}:J{fila_foto}", "VISTA PREVIA · FOTO ORIGINAL")
            ws.merge_range(f"E{caja_ini}:J{caja_fin}", "Aquí aparecerá la foto original",
                           f(bg_color="#FBF8F2", border=1, border_color="#E6DCCB", align="center", italic=True,
                             font_color="#B0A595"))
            r = caja_fin + 1
            self.seccion(ws, f"E{r}:J{r}", "ACCIONES · Excel de escritorio")
            b = lambda fila_x: f"E{fila_x}"  # noqa: E731
            h = lambda fila_x: f"H{fila_x}"  # noqa: E731
            self.boton("PROYECTO", b(r + 1), "1 · CARGAR FOTO", "CargarFoto", TEAL)
            self.boton("PROYECTO", h(r + 1), "NUEVO PROYECTO", "NuevoProyecto", NEUTRO, dx=6)
            self.boton("PROYECTO", b(r + 3), "2 · GENERAR IMAGEN CON IA", "GenerarImagen", TERRACOTA, ancho=552,
                       alto=50)
            self.boton("PROYECTO", b(r + 5), "COPIAR PROMPT + GEMINI WEB", "CopiarPromptGemini", MOSTAZA,
                       texto_oscuro=True)
            self.boton("PROYECTO", h(r + 5), "COPIAR PROMPT + CHATGPT", "CopiarPromptChatGPT", MOSTAZA, dx=6,
                       texto_oscuro=True)
            self.boton("PROYECTO", b(r + 7), "IMPORTAR RESULTADO", "ImportarResultado", TEAL)
            self.boton("PROYECTO", h(r + 7), "VER RESULTADO", "IrAResultado", NEUTRO, dx=6)
            self.boton("PROYECTO", b(r + 9), "AJUSTES · CLAVE API", "IrAAjustes", NEUTRO)
            self.boton("PROYECTO", h(r + 9), "ABRIR CARPETA DE IMÁGENES", "AbrirCarpetaResultados", NEUTRO, dx=6)
            ws.merge_range(f"E{r + 11}:J{r + 12}",
                           "Los botones funcionan en Excel de escritorio con las macros habilitadas. "
                           "En el móvil usa la sección 7 (modo asistido).",
                           f(font_size=9, italic=True, font_color=GRIS, text_wrap=True, valign="top"))
        self.imprimir(ws, f"B1:J{max(fila, caja_fin + 13)}" if self.macros else f"B1:C{fila}")
        ws.protect("", PROTECCION)
        ws.set_selection("C10")

    def resultado(self):
        ws = self.ws["RESULTADO"]
        f = self.f
        ws.hide_gridlines(2)
        ws.set_tab_color(MOSTAZA)
        ws.set_column("A:A", 1.5)
        ws.set_column("B:G", 12.5)
        ws.set_column("H:H", 2.5)
        ws.set_column("I:N", 12.5)
        ws.set_column("O:O", 1.5)
        ws.set_default_row(22)
        self.barra_titulo(ws, "B1:N1", "RESULTADO · ANTES Y DESPUÉS")
        ws.write_url("B2", "internal:'PROYECTO'!C10", f(font_color=TEAL, bold=True, underline=1, font_size=10),
                     "← Volver a PROYECTO")
        ws.set_row(1, 36)
        self.boton("RESULTADO", "I2", "GENERAR OTRA VERSIÓN", "GenerarOtraVersion", TERRACOTA, ancho=270, alto=38,
                   dy=4)
        self.boton("RESULTADO", "L2", "ABRIR CARPETA", "AbrirCarpetaResultados", NEUTRO, ancho=270, alto=38,
                   dx=6, dy=4)
        ws.set_row(2, 10)
        self.seccion(ws, "B4:G4", "ANTES · foto original")
        self.seccion(ws, "I4:N4", "DESPUÉS · propuesta de reforma", TERRACOTA)
        hueco = f(bg_color="#FBF8F2", border=1, border_color="#E6DCCB", align="center", italic=True,
                  font_color="#B0A595", text_wrap=True)
        ws.merge_range("B5:G22", "La foto original aparecerá aquí" if self.macros else
                       "Inserta aquí la foto original (Insertar › Imagen)", hueco)
        ws.merge_range("I5:N22", "La imagen generada aparecerá aquí" if self.macros else
                       "Inserta aquí la imagen generada (Insertar › Imagen)", hueco)
        self.nombre("CajaAntes", "RESULTADO", "B5:G22")
        self.nombre("CajaDespues", "RESULTADO", "I5:N22")

        etq = f(bold=True, font_size=10, bg_color=ARENA, border=1, border_color="#E6DCCB", indent=1)
        val = f(font_size=10, border=1, border_color="#D9D9D9", text_wrap=True, valign="top")
        for k, (t, n) in enumerate([("Fecha", "ResFecha"), ("Generado con", "ResMotor"), ("Archivo", "ResArchivo")]):
            fila = 24 + k
            ws.write(f"B{fila}", t, etq)
            ws.merge_range(f"C{fila}:N{fila}", "", val)
            self.nombre(n, "RESULTADO", f"C{fila}")
        ws.write("B27", "Reforma", f(bold=True, font_size=10, bg_color=ARENA, border=1, border_color="#E6DCCB",
                                     indent=1, valign="top"))
        ws.merge_range("C27:N30", "", val)
        self.nombre("ResResumen", "RESULTADO", "C27")
        ws.merge_range("B32:N33",
                       "Visualización orientativa generada con inteligencia artificial: puede contener imprecisiones "
                       "en materiales o mobiliario y no sustituye a un proyecto técnico ni a un presupuesto.",
                       f(font_size=9, italic=True, font_color=GRIS, text_wrap=True, valign="top"))
        self.imprimir(ws, "B1:N33", horizontal=True)
        ws.protect("", PROTECCION)

    def historial(self):
        ws = self.ws["HISTORIAL"]
        f = self.f
        ws.hide_gridlines(2)
        ws.set_tab_color(NEUTRO)
        anchos = [("A:A", 1.5), ("B:B", 5), ("C:C", 16), ("D:D", 26), ("E:E", 20), ("F:F", 26), ("G:G", 16),
                  ("H:H", 22), ("I:I", 30), ("J:J", 9), ("K:K", 11), ("L:L", 34), ("M:M", 70)]
        celda = f(font_size=9, text_wrap=True, valign="top")
        for rango, ancho in anchos:
            ws.set_column(rango, ancho, celda if rango != "A:A" else None)
        self.barra_titulo(ws, "B1:M1", "HISTORIAL · imágenes generadas")
        ws.merge_range("B2:M2", "Se rellena solo cada vez que generas o importas una imagen desde Excel de escritorio. "
                                "Pulsa el nombre del archivo para abrir la imagen.",
                       f(font_size=9, italic=True, font_color=GRIS))
        cab = f(bold=True, font_size=10, font_color=BLANCO, bg_color=TEAL, text_wrap=True, border=1,
                border_color=TEAL)
        ws.write_row("B4", ["Nº", "Fecha", "Proyecto", "Estancia", "Tipo de reforma", "Estilo", "Proporción",
                            "Motor IA", "Seg.", "Resultado", "Imagen", "Resumen de la reforma"], cab)
        ws.set_row(3, 28)
        self.nombre("HistCabecera", "HISTORIAL", "B4")
        ws.freeze_panes(4, 0)
        self.imprimir(ws, "B1:M200", horizontal=True, paginas_alto=0)
        ws.repeat_rows(3)

    def ajustes(self):
        ws = self.ws["AJUSTES"]
        f = self.f
        ws.hide_gridlines(2)
        ws.set_tab_color(GRIS)
        ws.set_column("A:A", 1.5)
        ws.set_column("B:B", 26)
        ws.set_column("C:C", 40)
        ws.set_column("D:D", 2.5)
        ws.set_column("E:H", 12.5)
        ws.set_default_row(22)
        etiqueta = f(bold=True, font_size=10, bg_color=ARENA, border=1, border_color="#E6DCCB", indent=1)
        entrada = f(bg_color=CREMA, border=1, border_color=BORDE_CREMA, locked=False, indent=1)
        calculada = f(border=1, border_color="#D9D9D9", font_size=10)
        nota = f(font_size=9, italic=True, font_color=GRIS, text_wrap=True, valign="top")

        self.barra_titulo(ws, "B1:C1", "AJUSTES · Motor de IA")
        ws.merge_range("E1:H1", "", f(bg_color=TEAL))
        ws.write_url("B2", "internal:'PROYECTO'!C10", f(font_color=TEAL, bold=True, underline=1, font_size=10),
                     "← Volver a PROYECTO")
        ws.merge_range("B3:C3", "Solo para la generación automática en Excel de escritorio. En el móvil no hace falta.",
                       nota)

        self.seccion(ws, "B5:C5", "MOTOR DE IA")
        filas = [
            ("Motor", "CfgMotor", L.MOTOR[0][0], "Motor", True,
             "Gemini conserva muy bien la estructura. OpenAI es la alternativa."),
            ("Modelo Gemini", "CfgModeloGemini", L.MODELO_GEMINI[0][0], "ModeloGemini", False,
             "Puedes escribir otro nombre de modelo si Google publica uno nuevo."),
            ("Modelo OpenAI", "CfgModeloOpenAI", L.MODELO_OPENAI[0][0], "ModeloOpenAI", False,
             "Puedes escribir otro nombre de modelo si OpenAI publica uno nuevo."),
        ]
        for k, (t, n, v, lista, estricto, msg) in enumerate(filas):
            fila = 6 + k
            ws.write(f"B{fila}", t, etiqueta)
            ws.write(f"C{fila}", v, entrada)
            self.nombre(n, "AJUSTES", f"C{fila}")
            self.desplegable("AJUSTES", f"C{fila}", lista, t, msg, estricto)
        ws.write("B9", "Carpeta de imágenes", etiqueta)
        ws.write("C9", "", entrada)
        self.nombre("CfgCarpeta", "AJUSTES", "C9")
        ws.data_validation("C9", {"validate": "any", "input_title": "Carpeta de imágenes",
                                  "input_message": "Vacío = Imágenes\\Gaudi Reformas. O escribe una ruta completa, "
                                                   "p. ej. C:\\Obras\\Renders"})

        self.seccion(ws, "B11:C11", "CLAVES API · en este equipo, no en el archivo")
        ws.write("B12", "Clave Gemini", etiqueta)
        ws.write("C12", "No configurada en este equipo", calculada)
        self.nombre("EstadoClaveGemini", "AJUSTES", "C12")
        ws.write("B13", "Clave OpenAI", etiqueta)
        ws.write("C13", "No configurada en este equipo", calculada)
        self.nombre("EstadoClaveOpenAI", "AJUSTES", "C13")
        ws.merge_range("B14:C14", "Pulsa GUARDAR CLAVE, pega tu clave y después PROBAR CONEXIÓN. Así puedes "
                                  "compartir este Excel sin compartir tu clave.", nota)
        ws.set_row(13, 30)

        self.seccion(ws, "B16:C16", "DÓNDE CONSEGUIR LA CLAVE")
        link = f(font_color=TEAL, underline=1, border=1, border_color="#D9D9D9", font_size=10)
        ws.write("B17", "Google Gemini", etiqueta)
        ws.write_url("C17", "https://aistudio.google.com/apikey", link, "aistudio.google.com/apikey")
        ws.write("B18", "OpenAI", etiqueta)
        ws.write_url("C18", "https://platform.openai.com/api-keys", link, "platform.openai.com/api-keys")
        ws.merge_range("B19:C19", "El coste por imagen depende del proveedor y del modelo: consulta sus precios. "
                                  "El modo asistido (sección 7 de PROYECTO) usa tu cuenta de Gemini o ChatGPT.", nota)
        ws.set_row(18, 30)

        self.seccion(ws, "B21:C21", "MODELOS DISPONIBLES (octubre 2026)")
        modelos = [
            ("gemini-3.1-flash-image", "Gemini · rápido y de gran calidad. Recomendado."),
            ("gemini-3-pro-image-preview", "Gemini · máximo detalle; más lento y caro."),
            ("gpt-image-1.5", "OpenAI · edición con alta fidelidad a la foto."),
            ("gpt-image-2", "OpenAI · última generación."),
        ]
        for k, (m, d) in enumerate(modelos):
            ws.write(f"B{22 + k}", m, f(font_size=9, font_name="Consolas", border=1, border_color="#E6DCCB",
                                        bg_color=ARENA))
            ws.write(f"C{22 + k}", d, f(font_size=9, border=1, border_color="#D9D9D9"))
        ws.merge_range("B26:C26", "Si un proveedor retira un modelo, escribe el nombre nuevo en la celda del modelo "
                                  "(se admite cualquier valor).", nota)
        ws.set_row(25, 30)

        if self.macros:
            self.seccion(ws, "E5:H5", "ACCIONES")
            self.boton("AJUSTES", "E6", "GUARDAR CLAVE GEMINI", "ConfigurarClaveGemini", TEAL, ancho=180, alto=36)
            self.boton("AJUSTES", "G6", "GUARDAR CLAVE OPENAI", "ConfigurarClaveOpenAI", TEAL, ancho=180, alto=36,
                       dx=4)
            self.boton("AJUSTES", "E8", "PROBAR CONEXIÓN", "ProbarConexion", TERRACOTA, ancho=180, alto=36)
            self.boton("AJUSTES", "G8", "BORRAR CLAVES", "BorrarClaves", NEUTRO, ancho=180, alto=36, dx=4)
            self.boton("AJUSTES", "E10", "ABRIR CARPETA", "AbrirCarpetaResultados", NEUTRO, ancho=180, alto=36)
            self.boton("AJUSTES", "G10", "VOLVER A PROYECTO", "IrAProyecto", NEUTRO, ancho=180, alto=36, dx=4)
        self.imprimir(ws, "B1:H26", horizontal=True)
        ws.protect("", PROTECCION)

    def construir(self):
        self.wb.set_properties({
            "title": "Gaudí · Visualizador de reformas",
            "subject": "Visualización de reformas con IA manteniendo la estructura original",
            "keywords": "reforma, interiorismo, IA, render, antes y después",
            "comments": "Generado con src/build_app.py",
        })
        self.listas()
        self.calc()
        self.inicio()
        self.proyecto()
        self.resultado()
        self.historial()
        self.ajustes()
        if self.macros:
            self.wb.add_vba_project(io.BytesIO(proyecto_vba()), is_stream=True)
        self.wb.close()


# =============================================================================
# Post-proceso: asigna las macros a las formas (botones redondeados)
# =============================================================================

def asignar_macros(ruta: Path) -> int:
    tmp = ruta.with_suffix(".tmp")
    total = 0
    with zipfile.ZipFile(ruta) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/drawings/drawing") and item.filename.endswith(".xml"):
                xml = data.decode("utf-8")

                def cambiar(m: re.Match) -> str:
                    nonlocal total
                    bloque = m.group(0)
                    d = re.search(r'descr="btn:(\w+)"', bloque)
                    if not d:
                        return bloque
                    total += 1
                    macro = d.group(1)
                    bloque = bloque.replace('<xdr:sp macro="" ', f'<xdr:sp macro="[0]!{macro}" ', 1)
                    bloque = bloque.replace(f'descr="btn:{macro}"', f'descr="Botón: {macro}"')
                    bloque = bloque.replace('<a:prstGeom prst="rect">', '<a:prstGeom prst="roundRect">', 1)
                    bloque = bloque.replace(' txBox="1"', "")
                    return bloque

                xml = re.sub(r'<xdr:sp macro="".*?</xdr:sp>', cambiar, xml, flags=re.S)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(ruta)
    return total


# =============================================================================
# Valores en caché (para que los visores sin cálculo muestren el prompt)
# =============================================================================

def calcular_cache() -> dict[str, object]:
    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory() as d:
        borrador = Path(d) / "borrador.xlsx"
        App(borrador, macros=False, cache={}).construir()
        # Perfil temporal de LibreOffice que fuerza el recálculo de las fórmulas al abrir
        perfil = Path(d) / "perfil"
        (perfil / "user").mkdir(parents=True)
        (perfil / "user" / "registrymodifications.xcu").write_text(
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
            'xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
            '<prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>'
            '</oor:items>')
        out = Path(d) / "out"
        subprocess.run(["soffice", f"-env:UserInstallation={perfil.as_uri()}", "--headless", "--calc",
                        "--convert-to", "xlsx", "--outdir", str(out), str(borrador)],
                       check=True, capture_output=True, timeout=180)
        wb = load_workbook(out / "borrador.xlsx", data_only=True)
        cache: dict[str, object] = {}
        for hoja in ("CALC", "PROYECTO"):
            for fila in wb[hoja].iter_rows():
                for celda in fila:
                    if celda.value is not None:
                        cache[f"{hoja}!{celda.coordinate}"] = celda.value
        return cache


def main() -> None:
    DIST.mkdir(exist_ok=True)
    cache = calcular_cache()
    App(XLSM, macros=True, cache=cache).construir()
    n = asignar_macros(XLSM)
    App(XLSX_MOVIL, macros=False, cache=cache).construir()
    print(f"OK  {XLSM.relative_to(ROOT)}  ({XLSM.stat().st_size // 1024} KB, {n} botones con macro)")
    print(f"OK  {XLSX_MOVIL.relative_to(ROOT)}  ({XLSX_MOVIL.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    sys.path.insert(0, str(SRC))
    main()
